"""Extract per-slide text from a PDF or PPTX and flag what each slide is.

Usage: python extract_slides.py <slides.pdf|slides.pptx> [--json]

Flags:
  FILLER   title / break / joke / agenda slide, safe to skip
  BUILD    mostly repeats the previous slide (animation step), merge with it
  EXERCISE in-class question or "what does this print" slide
  CODE     contains source code
  VISUAL   very little text, so the content is probably a diagram or image; look at it
"""
import json
import re
import sys
from difflib import SequenceMatcher

FILLER_WORDS = re.compile(
    r"\b(break time|joke|agenda|outline|announcements?|logistics|questions\?|"
    r"we.ll start again|thank you|office hours)\b", re.I)
EXERCISE_WORDS = re.compile(
    r"\b(exercise|question|quiz|what (does|would|is)|which of|poll|predict)\b", re.I)
CODE_LINE = re.compile(r"[;{}]\s*$|^\s*(def|class|public|private|return|for|while|if|#include|int|void)\b")


def read_pdf(path):
    from pypdf import PdfReader
    return [(p.extract_text() or "") for p in PdfReader(path).pages]


def read_pptx(path):
    from pptx import Presentation
    slides = []
    for s in Presentation(path).slides:
        parts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if s.has_notes_slide:
            parts.append("[speaker notes] " + s.notes_slide.notes_text_frame.text)
        slides.append("\n".join(parts))
    return slides


def clean(text):
    text = text.replace("", "-").replace("", "-").strip()
    lines = [l.rstrip() for l in text.splitlines() if l.strip()]
    # PowerPoint PDFs put the page number on its own line or glue it onto the title ("Polymorphism2")
    while lines and lines[0].strip().isdigit():
        lines.pop(0)
    if lines:
        lines[0] = re.sub(r"(?<=[A-Za-z!?)\"”])\d{1,3}$", "", lines[0])
    return lines


def classify(lines, prev_text):
    text = "\n".join(lines)
    flags = []
    code_lines = sum(bool(CODE_LINE.search(l)) for l in lines)
    if FILLER_WORDS.search(text) or (len(text) < 60 and code_lines == 0 and len(lines) <= 2):
        flags.append("FILLER")
    if prev_text and len(text) > 80 and SequenceMatcher(None, prev_text, text).ratio() > 0.8:
        flags.append("BUILD")
    if EXERCISE_WORDS.search(text):
        flags.append("EXERCISE")
    if code_lines >= 2:
        flags.append("CODE")
    if len(text) < 120 and "FILLER" not in flags:
        flags.append("VISUAL")
    return flags


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    raw = read_pptx(path) if path.lower().endswith(".pptx") else read_pdf(path)

    slides, prev = [], ""
    for i, text in enumerate(raw, 1):
        lines = clean(text)
        joined = "\n".join(lines)
        # Skip code lines when picking a title; exercise slides often start with code
        title = next((l for l in lines if not CODE_LINE.search(l)), lines[0] if lines else "")
        slides.append({"slide": i, "title": title,
                       "flags": classify(lines, prev), "text": joined})
        prev = joined

    if "--json" in sys.argv:
        print(json.dumps(slides, ensure_ascii=False, indent=1))
        return
    for s in slides:
        print(f"--- slide {s['slide']}: {s['title']}  [{', '.join(s['flags'])}]")
        print(s["text"])
    skip = [s["slide"] for s in slides if {"FILLER", "BUILD"} & set(s["flags"])]
    print(f"\n== {len(slides)} slides; likely skippable: {skip}")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
